from fastapi import APIRouter
from fastapi import Depends
from fastapi import HTTPException

from fastapi.responses import FileResponse

from sqlalchemy.orm import Session

from openpyxl import Workbook

from datetime import datetime
from datetime import timedelta

from typing import List

import os

from app.database import SessionLocal

from app.models.attendance import Attendance
from app.models.task import Task
from app.models.employees import Employee

from app.schemas.report import (
    ReportDateResponse
)

from app.utils.dependencies import (
    get_current_admin
)

router = APIRouter(
    prefix="/reports",
    tags=["Reports"]
)

def get_db():
    db = SessionLocal()

    try:
        yield db

    finally:
        db.close()

REPORTS_DIR = "reports"
os.makedirs(
    REPORTS_DIR,
    exist_ok=True
)

@router.get( # Get list of available report dates
    "/dates",
    response_model=List[ReportDateResponse]
)
def get_report_dates(
    db: Session = Depends(get_db),
    admin_id: int = Depends(get_current_admin)
):

    dates = set()

    attendance_dates = (
        db.query(
            Attendance.date
        )
        .distinct()
        .all()
    )

    for row in attendance_dates:

        if row[0]:

            dates.add(
                row[0].strftime(
                    "%d-%m-%Y"
                )
            )

    task_dates = (
        db.query(
            Task.assigned_at
        )
        .all()
    )

    for row in task_dates:

        if row[0]:

            dates.add(
                row[0].strftime(
                    "%d-%m-%Y"
                )
            )

    sorted_dates = sorted(
        dates,
        key=lambda date: datetime.strptime(
            date,
            "%d-%m-%Y"
        ),
        reverse=True
    )

    return [
        {"date": date}
        for date in sorted_dates
    ]

@router.get(   # Generate and download attendance report for a specific date
    "/attendance/{report_date}"
)
def attendance_report(
    report_date: str,
    db: Session = Depends(get_db),
    admin_id: int = Depends(get_current_admin)
):

    try:

        selected_date = datetime.strptime(
            report_date,
            "%d-%m-%Y"
        ).date()

    except ValueError:

        raise HTTPException(
            status_code=400,
            detail="Date format must be DD-MM-YYYY"
        )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Attendance"

    sheet.append([
        "Employee ID",
        "Employee Name",
        "Date",
        "Status",
        "Latitude",
        "Longitude"
    ])

    today = datetime.now().date()

    employees = (
        db.query(Employee)
        .filter(
            Employee.status == "active"
        )
        .all()
    )

    for employee in employees:

        attendance = (
            db.query(Attendance)
            .filter(
                Attendance.employee_id
                == employee.id,

                Attendance.date
                == selected_date
            )
            .first()
        )

        if attendance:

            final_status = (
                attendance.status
            )

            latitude = (
                attendance.latitude
            )

            longitude = (
                attendance.longitude
            )

        else:

            if selected_date < today:

                final_status = (
                    "absent"
                )

            else:

                final_status = (
                    "not_marked"
                )

            latitude = ""
            longitude = ""

        sheet.append([
            employee.id,
            employee.name,
            str(selected_date),
            final_status,
            latitude,
            longitude
        ])

    filename = os.path.join(
        REPORTS_DIR,
        f"attendance_report_{report_date}.xlsx"
    )

    workbook.save(filename)

    return FileResponse(
        path=filename,
        filename=os.path.basename(
            filename
        ),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@router.get(   # Generate and download task report for a specific date
    "/tasks/{report_date}"
)
def task_report(
    report_date: str,
    db: Session = Depends(get_db),
    admin_id: int = Depends(get_current_admin)
):

    try:

        selected_date = datetime.strptime(
            report_date,
            "%d-%m-%Y"
        ).date()

    except ValueError:

        raise HTTPException(
            status_code=400,
            detail="Date format must be DD-MM-YYYY"
        )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Tasks"

    sheet.append([
        "Task ID",
        "Title",
        "Description",
        "Employee ID",
        "Employee Name",
        "Status",
        "Is Late",
        "Assigned At",
        "Completed At"
    ])

    tasks = (
        db.query(Task)
        .all()
    )

    for task in tasks:

        if task.assigned_at.date() != selected_date:
            continue

        employee = (
            db.query(Employee)
            .filter(
                Employee.id == task.employee_id
            )
            .first()
        )

        employee_name = (
            employee.name
            if employee
            else "Unknown"
        )

        sheet.append([
        task.id,
        task.title,
        task.description,
        task.employee_id,
        employee_name,
        task.status,
        task.is_late,

        task.assigned_at.strftime(
            "%d-%m-%Y %H:%M:%S"
        ),

        task.completed_at.strftime(
            "%d-%m-%Y %H:%M:%S"
        )
        if task.completed_at
        else "Not Completed"
    ])

    filename = os.path.join(
        REPORTS_DIR,
        f"task_report_{report_date}.xlsx"
    )

    workbook.save(filename)

    return FileResponse(
        path=filename,
        filename=os.path.basename(filename),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get(     # Generate and download summary report for a specific date
    "/summary/{report_date}"
)
def summary_report(
    report_date: str,
    db: Session = Depends(get_db),
    admin_id: int = Depends(get_current_admin)
):

    try:

        selected_date = datetime.strptime(
            report_date,
            "%d-%m-%Y"
        ).date()

    except ValueError:

        raise HTTPException(
            status_code=400,
            detail="Date format must be DD-MM-YYYY"
        )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Summary"

    total_employees = (
        db.query(Employee)
        .filter(
            Employee.status == "active"
        )
        .count()
    )

    present = (
        db.query(Attendance)
        .filter(
            Attendance.date == selected_date,
            Attendance.status == "present"
        )
        .count()
    )

    late = (
        db.query(Attendance)
        .filter(
            Attendance.date == selected_date,
            Attendance.status == "late"
        )
        .count()
    )

    actual_absent = (
        db.query(Attendance)
        .filter(
            Attendance.date == selected_date,
            Attendance.status == "absent"
        )
        .count()
    )

    marked = (
        db.query(Attendance)
        .filter(
            Attendance.date == selected_date
        )
        .count()
    )

    missing = (
        total_employees - marked
    )
    today = datetime.now().date()
    if selected_date < today:
        absent = (
            actual_absent + missing
        )
        not_marked = 0

    else:
        absent = actual_absent
        not_marked = missing

    completed_tasks = 0

    tasks = (
        db.query(Task)
        .filter(
            Task.completed_at.isnot(None)
        )
        .all()
    )

    for task in tasks:

        if task.completed_at.date() == selected_date:
            completed_tasks += 1

    pending_tasks = (
        db.query(Task)
        .filter(
            Task.status == "pending"
        )
        .count()
    )

    sheet.append([
        "Metric",
        "Value"
    ])

    sheet.append([
        "Total Employees",
        total_employees
    ])

    sheet.append([
        "Present",
        present
    ])

    sheet.append([
        "Late",
        late
    ])

    sheet.append([
        "Absent",
        absent
    ])

    sheet.append([
        "Not Marked",
        not_marked
    ])

    sheet.append([
        "Pending Tasks",
        pending_tasks
    ])

    sheet.append([
        "Completed Tasks",
        completed_tasks
    ])

    filename = os.path.join(
        REPORTS_DIR,
        f"summary_report_{report_date}.xlsx"
    )

    workbook.save(filename)

    return FileResponse(
        path=filename,
        filename=os.path.basename(filename),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )